import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FUNCTIONAL_RESULTS = ROOT / "LMS_Test_Log_Results.json"
NFR_RESULTS = ROOT / "LMS_Non_Functional_Test_Results.json"
COMBINED_RESULTS = ROOT / "LMS_Test_Log_Results_NFR_Updated.json"
CHART_DIR = ROOT / "reports" / "nfr_updated_charts"
CHART_DIR.mkdir(parents=True, exist_ok=True)
TESTER = "Yong Di Lun"

NFR_REQUIREMENTS = [
    ("REQ_Q001", "The LMS shall maintain 100% transactional accuracy when handling concurrent book reservations to strictly prevent double-booking of a single physical copy.", "High", TESTER, "Reliability & Accuracy"),
    ("REQ_Q002", "The database shall ensure data integrity using foreign key constraints between User, Book, and Reservation tables to prevent orphaned records.", "Medium", TESTER, "Reliability & Accuracy"),
    ("REQ_Q003", "User passwords shall be securely hashed and salted (e.g., using bcrypt) prior to database storage; plain text passwords shall never be saved.", "High", TESTER, "Security"),
    ("REQ_Q004", "All client-to-server interactions shall be encrypted using HTTPS/TLS 1.2 or higher to protect login credentials and personal profile data during transit.", "High", TESTER, "Security"),
]

NFR_CASES = [
    {
        "id": "TC-Q001",
        "feature": "REQ_Q001",
        "objective": "Verify concurrent reservation transactional accuracy",
        "coverage": "TCOV-Q001",
        "input": "Two authenticated users reserve one book with Count = 1 at the same time.",
        "expected": "Only one reservation is created and the book count becomes 0. The second reservation must be rejected.",
        "requirements": "Prepare two dummy users and one available book with a single physical copy.",
        "dependency": "None",
    },
    {
        "id": "TC-Q002",
        "feature": "REQ_Q002",
        "objective": "Verify database foreign key integrity for Reservation records",
        "coverage": "TCOV-Q002",
        "input": "Inspect reserve table foreign keys and attempt to insert Reservation with user_id = 999999.",
        "expected": "Foreign key constraints exist and the orphaned reservation insert is rejected.",
        "requirements": "MySQL lms database must be accessible.",
        "dependency": "None",
    },
    {
        "id": "TC-Q003",
        "feature": "REQ_Q003",
        "objective": "Verify secure password hashing and salting",
        "coverage": "TCOV-Q003",
        "input": "Create or inspect two dummy user accounts with different passwords.",
        "expected": "Stored passwords are not plain text and use unique secure salted hashes such as bcrypt.",
        "requirements": "Users table must be accessible for verification.",
        "dependency": "None",
    },
    {
        "id": "TC-Q004",
        "feature": "REQ_Q004",
        "objective": "Verify HTTPS/TLS protection for client-server traffic",
        "coverage": "TCOV-Q004",
        "input": "Access /signin using HTTP and HTTPS.",
        "expected": "The login page and sensitive interactions are served over HTTPS/TLS 1.2 or higher.",
        "requirements": "LMS server must be running.",
        "dependency": "None",
    },
]

NFR_PROCEDURES = [
    ("TP-Q-001", "Concurrent Reservation Accuracy", ["TC-Q001"], [
        "Create two valid dummy member accounts.",
        "Create one dummy book record with Count = 1.",
        "Authenticate both dummy users.",
        "Trigger reservation for both users at the same time. [TC-Q001]",
        "Check reservation table count and final book count.",
    ], ["Remove dummy users, reservation records, and book records."]),
    ("TP-Q-002", "Database Foreign Key Integrity", ["TC-Q002"], [
        "Inspect foreign key metadata for the reserve table.",
        "Attempt to insert a Reservation record with invalid user_id = 999999. [TC-Q002]",
        "Check whether the database rejects the orphaned record.",
    ], ["Delete any temporary orphan record if it was inserted."]),
    ("TP-Q-003", "Secure Password Storage", ["TC-Q003"], [
        "Create or inspect two dummy user accounts.",
        "Read the password values stored in the users table. [TC-Q003]",
        "Verify that stored values are not plain text, are bcrypt-like, and are unique per password.",
    ], ["Remove dummy password test accounts."]),
    ("TP-Q-004", "HTTPS/TLS Protection", ["TC-Q004"], [
        "Open the LMS login page using HTTP.",
        "Attempt to open the LMS login page using HTTPS. [TC-Q004]",
        "Verify that sensitive client-to-server interaction is protected by HTTPS/TLS 1.2 or higher.",
    ], ["No cleanup required."]),
]

FUNCTIONAL_SEVERITY = {
    "TC-02-002": "Mission Critical", "TC-05-004": "Mission Critical", "TC-06-005": "Mission Critical",
    "TC-10-001": "Mission Critical", "TC-10-009": "Mission Critical", "TC-11-002": "Mission Critical",
    "TC-11-008": "Mission Critical", "TC-12-001": "Mission Critical", "TC-12-007": "Mission Critical",
    "TC-12-010": "Mission Critical", "TC-04-003": "Major", "TC-04-004": "Major", "TC-07-002": "Major",
    "TC-07-003": "Major", "TC-07-006": "Major", "TC-08-003": "Major", "TC-08-004": "Major",
    "TC-08-005": "Major", "TC-10-004": "Major", "TC-10-006": "Major", "TC-11-003": "Major",
    "TC-12-002": "Major", "TC-12-005": "Major", "TC-05-003": "Minor", "TC-10-003": "Minor",
    "TC-10-008": "Minor", "TC-11-005": "Minor", "TC-11-006": "Minor",
}
NFR_SEVERITY = {"TC-Q002": "Major", "TC-Q003": "Mission Critical", "TC-Q004": "Mission Critical"}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text="", bold=False, size=8, align=None, fill=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    for idx, line in enumerate(str(text).split("\n")):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Calibri"
        run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_title(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)


def combine_results():
    functional = json.loads(FUNCTIONAL_RESULTS.read_text(encoding="utf-8"))
    nfr = json.loads(NFR_RESULTS.read_text(encoding="utf-8"))
    rows = functional["results"] + nfr["results"]
    summary = {
        "total": len(rows),
        "pass": sum(1 for row in rows if row["pass_fail"] == "Pass"),
        "fail": sum(1 for row in rows if row["pass_fail"] == "Fail"),
        "not_executed": sum(1 for row in rows if row["pass_fail"] == "Not Executed"),
    }
    combined = {
        "tester": TESTER,
        "start_time": functional["start_time"],
        "end_time": nfr["end_time"],
        "results": rows,
        "summary": summary,
    }
    COMBINED_RESULTS.write_text(json.dumps(combined, indent=2), encoding="utf-8")
    return combined


def generate_requirements_doc():
    doc = Document()
    add_title(doc, "Non-Functional Requirements")
    for heading, rows in [
        ("Reliability & Accuracy", [r for r in NFR_REQUIREMENTS if r[4] == "Reliability & Accuracy"]),
        ("Security", [r for r in NFR_REQUIREMENTS if r[4] == "Security"]),
    ]:
        add_title(doc, heading, level=2)
        doc.add_paragraph(f"The {heading.lower()} requirements for the LMS are as follows:")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, header in enumerate(["Requirement ID", "Description", "Priority", "Author"]):
            set_cell(table.cell(0, idx), header, bold=True, fill="D9D9D9", align=WD_ALIGN_PARAGRAPH.CENTER)
        for req_id, desc, priority, author, _ in rows:
            cells = table.add_row().cells
            for idx, value in enumerate([req_id, desc, priority, author]):
                set_cell(cells[idx], value)
        doc.add_paragraph("")
    doc.save(ROOT / "LMS_Non_Functional_Requirements.docx")


def add_case_table(doc, case):
    table = doc.add_table(rows=6, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, (label, value) in enumerate([
        ("Test Case ID", case["id"]),
        ("Related Requirement ID", case["feature"]),
        ("Objective", case["objective"]),
        ("Covered Test Coverage Items", case["coverage"]),
    ]):
        set_cell(table.cell(row_idx, 0), label, bold=True, size=10)
        merged = table.cell(row_idx, 1).merge(table.cell(row_idx, 3))
        set_cell(merged, value, size=10)
    for col_idx, header in enumerate(["Input", "Expected Result", "Special Procedural Requirements", "Intercase Dependency"]):
        set_cell(table.cell(4, col_idx), header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    for col_idx, value in enumerate([case["input"], case["expected"], case["requirements"], case["dependency"]]):
        set_cell(table.cell(5, col_idx), value, size=9)
    caption(doc, f'{case["id"]} {case["objective"]} Test Case')


def generate_test_case_doc():
    doc = Document(ROOT / "LMS_Test_Case_Tables_Multi_Input_Format.docx")
    doc.add_page_break()
    add_title(doc, "Non-Functional Requirement Test Cases", level=1)
    for case in NFR_CASES:
        add_case_table(doc, case)
        doc.add_paragraph("")
    doc.save(ROOT / "LMS_Test_Case_Tables_NFR_Updated.docx")


def add_procedure_table(doc, proc):
    pid, objective, cases, setup, wrap = proc
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Test Procedure ID", pid),
        ("Objective", objective),
        ("Test Cases To Be Executed", ", ".join(cases)),
        ("Set Up", "\n".join(f"{i}.    {step}" for i, step in enumerate(setup, 1))),
        ("Wrap Up", "\n".join(f"{i}.    {step}" for i, step in enumerate(wrap, 1))),
    ]
    for idx, (label, value) in enumerate(rows):
        set_cell(table.cell(idx, 0), label, bold=True, size=10)
        set_cell(table.cell(idx, 1), value, size=10)
    caption(doc, f"{pid} {objective} Test Procedure")


def generate_procedure_doc():
    doc = Document(ROOT / "LMS_Test_Procedure_Specification.docx")
    doc.add_page_break()
    add_title(doc, "2.3.14 Non-Functional Requirement Test Procedure", level=2)
    doc.add_paragraph("Prior to execution of the following non-functional test procedures, these special requirements must be prepared:")
    doc.add_paragraph("i.    LMS application must be running against the test MySQL database.\nii.   Dummy users and book records must be available.\niii.  HTTPS verification must be performed against the deployed server endpoint.")
    for proc in NFR_PROCEDURES:
        add_procedure_table(doc, proc)
        doc.add_paragraph("")
    doc.save(ROOT / "LMS_Test_Procedure_Specification_NFR_Updated.docx")


def generate_log_doc(combined):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    title = add_title(doc, "Test Log")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = doc.add_table(rows=6, cols=4)
    info.style = "Table Grid"
    set_cell(info.cell(0, 0).merge(info.cell(0, 3)), "General Information", bold=True, fill="A6A6A6")
    set_cell(info.cell(1, 0), "Test Log Scope", fill="D9D9D9")
    set_cell(info.cell(1, 1).merge(info.cell(1, 3)), "This Test Log covers LMS functional and non-functional requirements, including REQ_Q001 to REQ_Q004.")
    set_cell(info.cell(2, 0), "Version Author", fill="D9D9D9")
    set_cell(info.cell(2, 1), TESTER)
    set_cell(info.cell(2, 2), "Contact Number", fill="D9D9D9")
    set_cell(info.cell(2, 3), "-")
    set_cell(info.cell(3, 0).merge(info.cell(3, 3)), "Activities Execution Information", bold=True, fill="A6A6A6")
    set_cell(info.cell(4, 0), "Tester Name", fill="D9D9D9")
    set_cell(info.cell(4, 1), TESTER)
    set_cell(info.cell(4, 2), "Execution Date", fill="D9D9D9")
    set_cell(info.cell(4, 3), datetime.fromisoformat(combined["end_time"]).strftime("%d/%m/%Y"))
    s = combined["summary"]
    set_cell(info.cell(5, 0), "Execution Summary", fill="D9D9D9")
    set_cell(info.cell(5, 1).merge(info.cell(5, 3)), f"Total: {s['total']}; Pass: {s['pass']}; Fail: {s['fail']}; Not Executed: {s['not_executed']}")
    doc.add_paragraph("")
    table = doc.add_table(rows=2, cols=9)
    table.style = "Table Grid"
    set_cell(table.cell(0, 0).merge(table.cell(0, 8)), "Procedure Result", bold=True, fill="A6A6A6")
    headers = ["Requirement ID", "Test Design ID", "Test Case ID", "Test Procedure ID", "Type of Testing", "Tool", "Pass/Fail", "Incident ID", "Remark"]
    for idx, h in enumerate(headers):
        set_cell(table.cell(1, idx), h, bold=True, fill="D9D9D9", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in combined["results"]:
        cells = table.add_row().cells
        vals = [row["requirement_id"], row["test_design_id"], row["test_case_id"], row["test_procedure_id"], row["type_of_testing"], row["tool"], row["pass_fail"], row["incident_id"], row.get("actual") or row.get("remark", "-")]
        for idx, val in enumerate(vals):
            set_cell(cells[idx], val, size=7)
    doc.add_paragraph("")
    add_environment_section(doc)
    doc.add_paragraph("")
    add_anomalous_event_section(doc, combined)
    doc.save(ROOT / "LMS_Test_Log_Yong_Di_Lun_NFR_Updated.docx")


def anomaly_text(row):
    if row.get("actual"):
        return f"{row.get('remark', 'Incident recorded.')} Actual result: {row['actual']}"
    return row.get("remark", "-")


def add_environment_section(doc):
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell(table.cell(0, 0).merge(table.cell(0, 1)), "Environment Information", bold=True, fill="A6A6A6")
    set_cell(table.cell(1, 0), "Requested Test Environment", bold=True, fill="D9D9D9", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table.cell(1, 1), "Test Environment After Changes", bold=True, fill="D9D9D9", align=WD_ALIGN_PARAGRAPH.CENTER)
    requested = (
        "Hardware: Local Windows workstation\n"
        "Software: Python 3.13.5, Selenium 4, Chrome headless, Playwright browser snapshots, "
        "Python Requests/Postman-compatible HTTP automation, Python chart/report generation\n"
        "Application URL: http://127.0.0.1:5000\n"
        "Database: MySQL lms database\n"
        "Reference: LMS_TPS_1.0.0, LMS_TL_1.0.0, LMS_TSR_1.0.0\n"
        "Included NFR Scope: REQ_Q001, REQ_Q002, REQ_Q003, REQ_Q004"
    )
    after = (
        "No application logic was changed. Temporary functional and non-functional test records were "
        "created and cleaned up by the test runner. Controlled database-failure simulations were "
        "performed in the test harness only. NFR checks were added for concurrent reservation accuracy, "
        "foreign key integrity, password hashing/salting, and HTTPS/TLS protection."
    )
    set_cell(table.cell(2, 0), requested)
    set_cell(table.cell(2, 1), after)


def add_anomalous_event_section(doc, combined):
    failures = [row for row in combined["results"] if row["pass_fail"] == "Fail"]
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell(table.cell(0, 0).merge(table.cell(0, 2)), "Anomalous Event", bold=True, fill="A6A6A6")
    for idx, header in enumerate(["Unexpected event occurred", "Test Procedure ID", "Anomaly Reporter Name"]):
        set_cell(table.cell(1, idx), header, bold=True, fill="D9D9D9", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in failures:
        cells = table.add_row().cells
        set_cell(cells[0], anomaly_text(row), size=7)
        set_cell(cells[1], row["test_procedure_id"], size=7)
        set_cell(cells[2], TESTER, size=7)


def generate_environment_anomaly_markdown(combined):
    failures = [row for row in combined["results"] if row["pass_fail"] == "Fail"]
    lines = [
        "# Updated Test Log Sections",
        "",
        "## Environment Information",
        "",
        "| Requested Test Environment | Test Environment After Changes |",
        "|---|---|",
        "| Hardware: Local Windows workstation<br>Software: Python 3.13.5, Selenium 4, Chrome headless, Playwright browser snapshots, Python Requests/Postman-compatible HTTP automation, Python chart/report generation<br>Application URL: http://127.0.0.1:5000<br>Database: MySQL lms database<br>Reference: LMS_TPS_1.0.0, LMS_TL_1.0.0, LMS_TSR_1.0.0<br>Included NFR Scope: REQ_Q001, REQ_Q002, REQ_Q003, REQ_Q004 | No application logic was changed. Temporary functional and non-functional test records were created and cleaned up by the test runner. Controlled database-failure simulations were performed in the test harness only. NFR checks were added for concurrent reservation accuracy, foreign key integrity, password hashing/salting, and HTTPS/TLS protection. |",
        "",
        "## Anomalous Event",
        "",
        "| Unexpected event occurred | Test Procedure ID | Anomaly Reporter Name |",
        "|---|---|---|",
    ]
    for row in failures:
        message = anomaly_text(row).replace("|", "\\|")
        lines.append(f"| {message} | {row['test_procedure_id']} | {TESTER} |")
    (ROOT / "LMS_Test_Log_Environment_Anomalous_NFR_Updated.md").write_text("\n".join(lines), encoding="utf-8")


def draw_bar_chart(path, labels, values, colors, title):
    img = Image.new("RGB", (900, 520), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    draw.text((30, 20), title, fill="black", font=font)
    max_value = max(values) if values else 1
    x0, y0, width, height = 70, 410, 620, 310
    draw.line((x0, y0, x0 + width, y0), fill="black")
    draw.line((x0, y0, x0, y0 - height), fill="black")
    bar_w = width // (len(values) * 2)
    for idx, (label, value, color) in enumerate(zip(labels, values, colors)):
        x = x0 + 35 + idx * bar_w * 2
        bar_h = int((value / max_value) * (height - 35))
        draw.rectangle((x, y0 - bar_h, x + bar_w, y0), fill=color)
        draw.text((x, y0 - bar_h - 18), str(value), fill="black", font=font)
        draw.text((x - 10, y0 + 15), label, fill="black", font=font)
    img.save(path)


def generate_summary_doc(combined):
    failed = [r for r in combined["results"] if r["pass_fail"] == "Fail"]
    severity = Counter({}) 
    for row in failed:
        severity[NFR_SEVERITY.get(row["test_case_id"], FUNCTIONAL_SEVERITY.get(row["test_case_id"], "Major"))] += 1
    s = combined["summary"]
    draw_bar_chart(CHART_DIR / "test_case_coverage_nfr_updated.png", ["Total", "Passed", "Failed", "Not Exec"], [s["total"], s["pass"], s["fail"], s["not_executed"]], ["black", "#5A92D6", "#B9CEE8", "#8566A8"], "Test Case Coverage - Functional + Non-Functional")
    draw_bar_chart(CHART_DIR / "defect_severity_nfr_updated.png", ["Critical", "Major", "Minor", "Total"], [severity["Mission Critical"], severity["Major"], severity["Minor"], s["fail"]], ["black", "#B9CEE8", "#A7C957", "#8566A8"], "Defects Found by Severity")
    doc = Document()
    title = add_title(doc, "Test Summary Report")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=5, cols=4)
    table.style = "Table Grid"
    set_cell(table.cell(0, 0).merge(table.cell(0, 3)), "General Information", bold=True, fill="A6A6A6")
    set_cell(table.cell(1, 0), "Test Report Description", fill="D9D9D9")
    set_cell(table.cell(1, 1).merge(table.cell(1, 3)), "This Test Summary Report includes LMS functional and non-functional test results.")
    set_cell(table.cell(2, 0), "Version Author", fill="D9D9D9")
    set_cell(table.cell(2, 1), TESTER)
    set_cell(table.cell(2, 2), "Contact Number", fill="D9D9D9")
    set_cell(table.cell(2, 3), "-")
    set_cell(table.cell(3, 0), "Associated test plan reference", fill="D9D9D9")
    set_cell(table.cell(3, 1).merge(table.cell(3, 3)), "LMS_TP_1.0.0")
    set_cell(table.cell(4, 0), "Included NFRs", fill="D9D9D9")
    set_cell(table.cell(4, 1).merge(table.cell(4, 3)), "REQ_Q001, REQ_Q002, REQ_Q003, REQ_Q004")
    doc.add_paragraph("")
    summary = doc.add_table(rows=3, cols=5)
    summary.style = "Table Grid"
    set_cell(summary.cell(0, 0).merge(summary.cell(0, 4)), "Test Status Report Summary", bold=True, fill="A6A6A6")
    for idx, h in enumerate(["Planned", "Remaining", "Completed", "Passed", "Failed"]):
        set_cell(summary.cell(1, idx), h, bold=True, fill="D9D9D9")
    for idx, v in enumerate([s["total"], 0, s["total"], s["pass"], s["fail"]]):
        set_cell(summary.cell(2, idx), str(v))
    doc.add_paragraph(f"Current Incident Status: New/Open = {s['fail']}, Rejected = 0, Resolved = 0, Deferred = 0.")
    doc.add_picture(str(CHART_DIR / "test_case_coverage_nfr_updated.png"), width=Inches(5.8))
    doc.add_picture(str(CHART_DIR / "defect_severity_nfr_updated.png"), width=Inches(5.8))
    doc.add_paragraph(f"Result Summary: {s['total']} total test cases were completed. {s['pass']} passed and {s['fail']} failed. The NFR update added 4 test cases, with 1 pass and 3 failures.")
    doc.add_paragraph("Conclusion: The current LMS version is not recommended for release until open functional and non-functional incidents are fixed, especially security and data-integrity defects.")
    doc.save(ROOT / "LMS_Test_Summary_Report_Yong_Di_Lun_NFR_Updated.docx")


def generate_incident_doc(combined):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    title = add_title(doc, "Test Incident Report")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    set_cell(table.cell(0, 0).merge(table.cell(0, 1)), "Incident report identifier (Table 1)", bold=True, fill="A6A6A6")
    set_cell(table.cell(1, 0), "Scope", fill="D9D9D9")
    set_cell(table.cell(1, 1), "This Test Incident Report covers all open functional and non-functional LMS incidents after adding REQ_Q001 to REQ_Q004.")
    set_cell(table.cell(2, 0), "References", fill="D9D9D9")
    set_cell(table.cell(2, 1), "LMS_TPS_1.0.0\nLMS_TL_1.0.0\nLMS_TSR_1.0.0")
    doc.add_paragraph("")
    failures = [r for r in combined["results"] if r["pass_fail"] == "Fail"]
    for row in failures:
        severity = NFR_SEVERITY.get(row["test_case_id"], FUNCTIONAL_SEVERITY.get(row["test_case_id"], "Major"))
        incident = doc.add_table(rows=12, cols=2)
        incident.style = "Table Grid"
        set_cell(incident.cell(0, 0), "Test Incident Number:", fill="D9D9D9")
        set_cell(incident.cell(0, 1), row["incident_id"])
        set_cell(incident.cell(1, 0), "Summary", fill="D9D9D9")
        set_cell(incident.cell(1, 1), row.get("objective") or row.get("remark", "Incident recorded during test execution."))
        set_cell(incident.cell(2, 0), "Context", fill="D9D9D9")
        set_cell(incident.cell(2, 1), f"{row['requirement_id']} / {row['test_case_id']} / {row['test_procedure_id']}")
        set_cell(incident.cell(3, 0), "Expected Result", fill="D9D9D9")
        set_cell(incident.cell(3, 1), row.get("expected", "Refer to test case specification."))
        set_cell(incident.cell(4, 0), "Actual Result", fill="D9D9D9")
        set_cell(incident.cell(4, 1), row.get("actual") or row.get("remark", "-"))
        set_cell(incident.cell(5, 0), "Procedure to reproduce", fill="D9D9D9")
        set_cell(incident.cell(5, 1), f"Execute {row['test_procedure_id']} and observe {row['test_case_id']}.")
        set_cell(incident.cell(6, 0), "Test Environment", fill="D9D9D9")
        set_cell(incident.cell(6, 1), "Local Windows workstation, Flask LMS server, MySQL lms database, Selenium/HTTP test harness.")
        set_cell(incident.cell(7, 0), "Attempt to repeat", fill="D9D9D9")
        set_cell(incident.cell(7, 1), "Incident reproduced during automated executable test run.")
        set_cell(incident.cell(8, 0), "Tester's Name", fill="D9D9D9")
        set_cell(incident.cell(8, 1), TESTER)
        set_cell(incident.cell(9, 0), "Status of Incident", fill="D9D9D9")
        set_cell(incident.cell(9, 1), "☒ Open    ☐ Assigned for Resolution    ☐ Retested with the fix confirmed    ☐ Fixed")
        set_cell(incident.cell(10, 0), "Impact / Priority", fill="D9D9D9")
        set_cell(incident.cell(10, 1), f"{severity}\nPriority: {'Immediate' if severity != 'Minor' else 'Deferred'}")
        set_cell(incident.cell(11, 0), "Corrective Action", fill="D9D9D9")
        set_cell(incident.cell(11, 1), "Developer should correct the failed requirement and rerun regression tests.")
        doc.add_paragraph("")
    doc.save(ROOT / "LMS_Test_Incident_Report_Yong_Di_Lun_NFR_Updated.docx")


def generate_markdown(combined):
    lines = [
        "# LMS NFR Updated Documents",
        "",
        f"Total test cases: {combined['summary']['total']}",
        f"Passed: {combined['summary']['pass']}",
        f"Failed: {combined['summary']['fail']}",
        "Non-functional requirements added: REQ_Q001, REQ_Q002, REQ_Q003, REQ_Q004",
        "",
        "| Test Case | Requirement | Result | Incident | Actual |",
        "|---|---|---|---|---|",
    ]
    for row in json.loads(NFR_RESULTS.read_text(encoding="utf-8"))["results"]:
        lines.append(f"| {row['test_case_id']} | {row['requirement_id']} | {row['pass_fail']} | {row['incident_id']} | {row['actual']} |")
    (ROOT / "LMS_NFR_Updated_Document_Index.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    combined = combine_results()
    generate_requirements_doc()
    generate_test_case_doc()
    generate_procedure_doc()
    generate_log_doc(combined)
    generate_summary_doc(combined)
    generate_incident_doc(combined)
    generate_markdown(combined)
    generate_environment_anomaly_markdown(combined)
    print(json.dumps(combined["summary"], indent=2))
    print("Generated NFR updated documents.")


if __name__ == "__main__":
    main()
