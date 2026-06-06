import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS_PATH = ROOT / "LMS_Test_Log_Results_NFR_Updated.json"
CASE_SPEC_PATH = ROOT / "LMS_Test_Case_Tables.md"
SNAPSHOT_DIR = ROOT / "reports" / "incident_snapshots"
OUT_DOCX = ROOT / "LMS_Test_Incident_Report_Yong_Di_Lun.docx"
FALLBACK_DOCX = ROOT / "LMS_Test_Incident_Report_Yong_Di_Lun_NFR_Updated.docx"
OUT_MD = ROOT / "LMS_Test_Incident_Report_Yong_Di_Lun.md"
TESTER = "Yong Di Lun"

SEVERITY_BY_CASE = {
    "TC-02-002": "Mission Critical",
    "TC-05-004": "Mission Critical",
    "TC-06-005": "Mission Critical",
    "TC-10-001": "Mission Critical",
    "TC-10-009": "Mission Critical",
    "TC-11-002": "Mission Critical",
    "TC-11-008": "Mission Critical",
    "TC-12-001": "Mission Critical",
    "TC-12-007": "Mission Critical",
    "TC-12-010": "Mission Critical",
    "TC-04-003": "Major",
    "TC-04-004": "Major",
    "TC-07-002": "Major",
    "TC-07-003": "Major",
    "TC-07-006": "Major",
    "TC-08-003": "Major",
    "TC-08-004": "Major",
    "TC-08-005": "Major",
    "TC-10-004": "Major",
    "TC-10-006": "Major",
    "TC-11-003": "Major",
    "TC-12-002": "Major",
    "TC-12-005": "Major",
    "TC-05-003": "Minor",
    "TC-10-003": "Minor",
    "TC-10-008": "Minor",
    "TC-11-005": "Minor",
    "TC-11-006": "Minor",
    "TC-Q002": "Major",
    "TC-Q003": "Mission Critical",
    "TC-Q004": "Mission Critical",
}

NFR_REQUIREMENTS = {
    "REQ_Q001": ("Reliability & Accuracy", "The LMS shall maintain 100% transactional accuracy when handling concurrent book reservations to strictly prevent double-booking of a single physical copy."),
    "REQ_Q002": ("Reliability & Accuracy", "The database shall ensure data integrity using foreign key constraints between User, Book, and Reservation tables to prevent orphaned records."),
    "REQ_Q003": ("Security", "User passwords shall be securely hashed and salted prior to database storage; plain text passwords shall never be saved."),
    "REQ_Q004": ("Security", "All client-to-server interactions shall be encrypted using HTTPS/TLS 1.2 or higher."),
}

NFR_CASE_SPECS = {
    "TC-Q001": {
        "objective": "Verify 100% transactional accuracy for concurrent book reservations",
        "input": "Two authenticated users reserve one book with Count = 1 at the same time.",
        "expected": "Only one reservation is created and the book count becomes 0. The second reservation must be rejected.",
    },
    "TC-Q002": {
        "objective": "Verify database foreign key integrity between User, Book, and Reservation tables",
        "input": "Inspect reserve table foreign keys and attempt to insert Reservation with user_id = 999999.",
        "expected": "Database rejects orphaned Reservation records and defines foreign key constraints for user_id and book_id.",
    },
    "TC-Q003": {
        "objective": "Verify secure password hashing and salting before database storage",
        "input": "Create or inspect two dummy user accounts and read the stored password values.",
        "expected": "Passwords are not stored in plain text and each password hash uses a secure salted algorithm such as bcrypt.",
    },
    "TC-Q004": {
        "objective": "Verify HTTPS/TLS 1.2 or higher is used for client-to-server interactions",
        "input": "Access /signin using HTTP and HTTPS.",
        "expected": "Login and profile traffic are served through HTTPS/TLS 1.2 or higher.",
    },
}

SNAPSHOT_BY_CASE = {
    "TC-02-002": "login_wrong_password.png",
    "TC-07-002": "search_empty_redirect.png",
    "TC-07-006": "search_empty_redirect.png",
    "TC-10-001": "admin_add_book_404.png",
    "TC-10-009": "admin_add_book_404.png",
    "TC-11-005": "admin_delete_no_confirm.png",
    "TC-11-006": "admin_delete_no_confirm.png",
}

FEATURE_NAME = {
    "REQ-F001": "Sign Up",
    "REQ-F002": "Login",
    "REQ-F003": "Logout",
    "REQ-F004": "Manage Profile",
    "REQ-F005": "View Book List",
    "REQ-F006": "View Book Details",
    "REQ-F007": "Search Book",
    "REQ-F008": "Reserve Book",
    "REQ-F009": "View Reserved Books",
    "REQ-F0010": "Add Book",
    "REQ-F0011": "Delete Book",
    "REQ-F0012": "Edit Book",
    "REQ_Q001": "Reliability & Accuracy - Concurrent Reservation Accuracy",
    "REQ_Q002": "Reliability & Accuracy - Database Foreign Key Integrity",
    "REQ_Q003": "Security - Secure Password Storage",
    "REQ_Q004": "Security - HTTPS/TLS Protection",
}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text="", bold=False, size=8, align=None, fill=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    for idx, line in enumerate(str(text).split("\n")):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Calibri"
        run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_image_to_cell(cell, image_path, width=4.0):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width))


def checkbox(selected):
    return "[x]" if selected else "[ ]"


def parse_case_specs():
    if not CASE_SPEC_PATH.exists():
        return {}
    text = CASE_SPEC_PATH.read_text(encoding="utf-8")
    specs = {}
    current = None
    for line in text.splitlines():
        if line.startswith("| Test Case ID |"):
            current = line.strip("|").split("|")[1].strip()
            specs[current] = {}
        elif current and line.startswith("| Objective |"):
            specs[current]["objective"] = line.strip("|").split("|")[1].strip()
        elif current and line.startswith("| Input |") and "Expected Result" in line:
            continue
        elif current and line.startswith("| ") and not line.startswith("|---"):
            parts = [part.strip() for part in line.strip("|").split("|")]
            if len(parts) == 4 and parts[0] and parts[1] and "Test Case ID" not in parts[0]:
                specs[current]["input"] = parts[0]
                specs[current]["expected"] = parts[1]
                specs[current]["requirements"] = parts[2]
                specs[current]["dependency"] = parts[3]
    return specs


def corrective_action(row, severity):
    case_id = row["test_case_id"]
    feature = row["requirement_id"]
    if case_id == "TC-Q002":
        return "Add foreign key constraints between reservation records and the related user/book records. Remove orphaned records and rerun integrity tests."
    if case_id == "TC-Q003":
        return "Replace the current password storage method with a secure salted hashing algorithm such as bcrypt. Force password reset or migration for existing test accounts, then rerun authentication/security tests."
    if case_id == "TC-Q004":
        return "Deploy the LMS behind HTTPS/TLS 1.2 or higher and redirect sensitive HTTP routes to HTTPS. Rerun login/profile traffic verification."
    if case_id == "TC-02-002":
        return "Fix password verification so the supplied password is hashed and compared correctly. Add regression tests for wrong-password login."
    if feature == "REQ-F004":
        return "Add server-side validation for profile name and email fields before updating user records."
    if feature == "REQ-F005":
        return "Ensure book list displays required title, author, availability/count fields and catches database retrieval errors gracefully."
    if feature == "REQ-F006":
        return "Catch book detail retrieval exceptions and display a controlled user-facing error message instead of HTTP 500."
    if feature == "REQ-F007":
        return "Add search input validation for empty keyword, maximum length, and validation message display."
    if feature == "REQ-F008":
        return "Enforce reservation limit and duplicate reservation checks in backend logic and display clear reservation failure messages."
    if feature == "REQ-F0010":
        return "Implement the admin Add Book submit route correctly and add server-side validation for required fields, quantity limits, and duplicate names."
    if feature == "REQ-F0011":
        return "Add explicit delete confirmation, non-existing book handling, and database exception handling for admin deletion."
    if feature == "REQ-F0012":
        return "Implement admin Edit Book POST persistence and validation for invalid book ID, quantity constraints, and partial updates."
    if severity == "Mission Critical":
        return "Developer must fix the underlying system failure and rerun the related regression tests before release."
    return "Developer should correct the failed requirement and rerun the related regression tests."


def actual_result(row):
    if row.get("actual"):
        return row["actual"]
    text = row["remark"]
    if "HTTP 500" in text:
        return f"{text} The system returned an unhandled server error."
    if "404" in text:
        return f"{text} The system navigated to a missing route."
    return text


def unexpected_outcome(row):
    severity = SEVERITY_BY_CASE.get(row["test_case_id"], "Major")
    if row["test_case_id"] == "TC-Q002":
        return "Database integrity is not enforced, allowing orphan reservation data to be created."
    if row["test_case_id"] == "TC-Q003":
        return "Password values are not stored using a secure salted password-hashing scheme."
    if row["test_case_id"] == "TC-Q004":
        return "Sensitive client-to-server interactions are available over HTTP and HTTPS/TLS is not configured for the tested endpoint."
    if severity == "Mission Critical":
        return "System behavior prevents the requirement from being completed correctly or causes a server/system failure."
    if severity == "Major":
        return "System behavior does not follow the expected business rule or validation requirement."
    return "System display or procedural behavior does not match the expected specification."


def procedure_to_reproduce(row, spec):
    steps = [
        f"Execute test procedure {row['test_procedure_id']}.",
        f"Execute test case {row['test_case_id']}.",
    ]
    if spec.get("input"):
        steps.append(f"Apply test data/input: {spec['input']}.")
    steps.append("Observe the actual system result and compare it with the expected result.")
    return "\n".join(f"{idx}.    {step}" for idx, step in enumerate(steps, 1))


def impact_text(severity):
    return (
        f"{checkbox(severity == 'Mission Critical')} Mission Critical : Application will not function or system fails\n"
        f"{checkbox(severity == 'Major')} Major : Severe problems but possible to work around\n"
        f"{checkbox(severity == 'Minor')} Minor : Does not impact the functionality or usability of the process but is not according to requirements/design specifications"
    )


def priority_text(severity):
    immediate = severity in {"Mission Critical", "Major"}
    return (
        f"{checkbox(immediate)} Immediate : Must be fixed as soon as possible\n"
        f"{checkbox(False)} Delayed : System is unstable but incident must be fixed prior to next level of test or shipment\n"
        f"{checkbox(not immediate)} Deferred : Defect can be left in if necessary due to time or costs"
    )


def add_identifier(doc):
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell(table.cell(0, 0).merge(table.cell(0, 1)), "Incident report identifier (Table 1)", bold=True, fill="A6A6A6")
    set_cell(table.cell(1, 0), "Scope", fill="D9D9D9")
    set_cell(
        table.cell(1, 1),
        "The scope of this Test Incident Report is any incident that occurred based on Test Procedures executed for the Library Management System features F001 to F0012 and non-functional requirements REQ_Q001 to REQ_Q004.",
    )
    set_cell(table.cell(2, 0), "References", fill="D9D9D9")
    set_cell(table.cell(2, 1), "LMS_TPS_1.0.0\nLMS_TL_1.0.0\nLMS_TSR_1.0.0\nLMS_NFR_1.0.0")
    doc.add_paragraph("")


def add_nfr_coverage(doc, data):
    nfr_rows = [row for row in data["results"] if row["requirement_id"].startswith("REQ_Q")]
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header in enumerate(["Requirement ID", "Category", "Test Case ID", "Test Procedure ID", "Pass/Fail", "Incident ID"]):
        set_cell(table.cell(0, idx), header, bold=True, fill="D9D9D9", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in nfr_rows:
        category = NFR_REQUIREMENTS.get(row["requirement_id"], ("Non-Functional", ""))[0]
        cells = table.add_row().cells
        values = [row["requirement_id"], category, row["test_case_id"], row["test_procedure_id"], row["pass_fail"], row["incident_id"]]
        for idx, value in enumerate(values):
            set_cell(cells[idx], value)
    doc.add_paragraph("")


def add_incident(doc, row, spec, incident_date):
    severity = SEVERITY_BY_CASE.get(row["test_case_id"], "Major")
    feature = FEATURE_NAME.get(row["requirement_id"], row["requirement_id"])
    title = spec.get("objective", row["remark"])
    table = doc.add_table(rows=21, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for table_row in table.rows:
        table_row.cells[0].width = Inches(1.45)
        table_row.cells[1].width = Inches(9.7)

    set_cell(table.cell(0, 0), "Test Incident Number:", fill="D9D9D9")
    set_cell(table.cell(0, 1), row["incident_id"])
    set_cell(table.cell(1, 0), "Summary", fill="D9D9D9")
    set_cell(table.cell(1, 1), title)
    set_cell(table.cell(2, 0), "Date and Time Incident:", fill="D9D9D9")
    set_cell(table.cell(2, 1), incident_date)
    set_cell(table.cell(3, 0), "Context:", fill="D9D9D9")
    set_cell(table.cell(3, 1), f"Library Management System (LMS 1.0.0)\nFeature: {feature}\nExecuted using {row['tool']}")
    set_cell(table.cell(4, 0).merge(table.cell(4, 1)), "Description of Incident", bold=True, fill="A6A6A6")
    set_cell(table.cell(5, 0), "Test\nProcedure:", fill="D9D9D9")
    set_cell(table.cell(5, 1), row["test_procedure_id"])
    set_cell(table.cell(6, 0), "Test Data:", fill="D9D9D9")
    set_cell(table.cell(6, 1), spec.get("input", "Refer to test case specification."))
    set_cell(table.cell(7, 0), "Expected\nResult:", fill="D9D9D9")
    set_cell(table.cell(7, 1), spec.get("expected", "System should meet the expected result defined in the test case specification."))
    set_cell(table.cell(8, 0), "Actual Result:", fill="D9D9D9")
    set_cell(table.cell(8, 1), actual_result(row))
    set_cell(table.cell(9, 0), "Unexpected\nOutcome:", fill="D9D9D9")
    set_cell(table.cell(9, 1), unexpected_outcome(row))
    snapshot_name = SNAPSHOT_BY_CASE.get(row["test_case_id"])
    if snapshot_name:
        snapshot_path = SNAPSHOT_DIR / snapshot_name
        if snapshot_path.exists():
            add_image_to_cell(table.cell(9, 1), snapshot_path, width=5.5)
    set_cell(table.cell(10, 0), "Procedure to\nreproduce the\nincident", fill="D9D9D9")
    set_cell(table.cell(10, 1), procedure_to_reproduce(row, spec))
    set_cell(table.cell(11, 0), "Test\nEnvironment", fill="D9D9D9")
    set_cell(
        table.cell(11, 1),
        "Hardware: Local Windows workstation\nSoftware: Python 3.13.5, Selenium 4, Playwright browser evidence, Python Requests/Postman-compatible HTTP automation\nDatabase: MySQL lms database\nReference: LMS_TPS_1.0.0",
    )
    set_cell(table.cell(12, 0), "Attempt to\nrepeat", fill="D9D9D9")
    set_cell(table.cell(12, 1), "Incident was reproduced during the automated executable test run. Retest after corrective action is required.")
    set_cell(table.cell(13, 0), "Tester's Name", fill="D9D9D9")
    set_cell(table.cell(13, 1), TESTER)
    set_cell(table.cell(14, 0), "Observer's\nName\n(witness)", fill="D9D9D9")
    set_cell(table.cell(14, 1), "-")
    set_cell(table.cell(15, 0).merge(table.cell(15, 1)), "Status of Incident", bold=True, fill="A6A6A6")
    set_cell(table.cell(16, 0).merge(table.cell(16, 1)), f"{checkbox(True)} Open        {checkbox(False)} Assigned for Resolution        {checkbox(False)} Retested with the fix confirmed\n{checkbox(False)} Approved for Resolution        {checkbox(False)} Fixed")
    set_cell(table.cell(17, 0).merge(table.cell(17, 1)), "Impact", bold=True, fill="A6A6A6")
    set_cell(table.cell(18, 0).merge(table.cell(18, 1)), impact_text(severity))
    set_cell(table.cell(19, 0).merge(table.cell(19, 1)), "Priority", bold=True, fill="A6A6A6")
    set_cell(table.cell(20, 0).merge(table.cell(20, 1)), priority_text(severity) + "\n\nDescription of the corrective action\n" + corrective_action(row, severity))
    doc.add_paragraph("")


def add_conclusion(doc, failures):
    table = doc.add_table(rows=3, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell(table.cell(0, 0), "Conclusions and Recommendations (Table 3)", bold=True, fill="A6A6A6")
    set_cell(
        table.cell(1, 0),
        f"{len(failures)} incidents were recorded during this test iteration. All incidents are open. Fatal and major incidents must be corrected before release, followed by regression testing using the same test procedures.",
    )
    set_cell(table.cell(2, 0), "Recommendation: Fix the reported incidents, update the affected modules, and rerun the executable test system without changing the testing procedure.")
    doc.add_paragraph("")


def add_approvals(doc):
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell(table.cell(0, 0).merge(table.cell(0, 2)), "Approvals", bold=True, fill="A6A6A6")
    for idx, header in enumerate(["Name", "Job Title", "Signature"]):
        set_cell(table.cell(1, idx), header, bold=True, fill="D9D9D9")
    rows = [
        (TESTER, "Test Lead\nLMS Test Team", ""),
        ("Project Supervisor", "Reviewer\nLMS Project", ""),
    ]
    for idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data):
            set_cell(table.cell(idx, col_idx), value)


def build_markdown(data, specs):
    failures = [row for row in data["results"] if row["pass_fail"] == "Fail"]
    nfr_rows = [row for row in data["results"] if row["requirement_id"].startswith("REQ_Q")]
    lines = [
        "# Test Incident Report",
        "",
        f"Tester: {TESTER}",
        "Scope: Functional requirements F001 to F0012 and non-functional requirements REQ_Q001 to REQ_Q004.",
        "References: LMS_TPS_1.0.0, LMS_TL_1.0.0, LMS_TSR_1.0.0, LMS_NFR_1.0.0",
        "",
        "## Non-Functional Requirement Coverage",
        "",
        "| Requirement ID | Category | Test Case ID | Test Procedure ID | Pass/Fail | Incident ID |",
        "|---|---|---|---|---|---|",
    ]
    for row in nfr_rows:
        category = NFR_REQUIREMENTS.get(row["requirement_id"], ("Non-Functional", ""))[0]
        lines.append(f"| {row['requirement_id']} | {category} | {row['test_case_id']} | {row['test_procedure_id']} | {row['pass_fail']} | {row['incident_id']} |")
    lines.extend([
        "",
        "## Incidents",
        "",
    ])
    for row in failures:
        spec = specs.get(row["test_case_id"], {})
        severity = SEVERITY_BY_CASE.get(row["test_case_id"], "Major")
        lines.extend(
            [
                f"## {row['incident_id']} - {row['test_case_id']}",
                "",
                f"Summary: {spec.get('objective', row['remark'])}",
                f"Severity: {severity}",
                f"Status: Open",
                f"Test Procedure: {row['test_procedure_id']}",
                f"Expected Result: {spec.get('expected', '-')}",
                f"Actual Result: {actual_result(row)}",
                f"Corrective Action: {corrective_action(row, severity)}",
                "",
            ]
        )
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main():
    data = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    specs = parse_case_specs()
    specs.update(NFR_CASE_SPECS)
    failures = [row for row in data["results"] if row["pass_fail"] == "Fail"]
    incident_date = datetime.fromisoformat(data["end_time"]).strftime("%d/%m/%Y %H:%M")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(8)

    title = doc.add_heading("Test Incident Report", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_identifier(doc)
    add_nfr_coverage(doc, data)
    for row in failures:
        add_incident(doc, row, specs.get(row["test_case_id"], {}), incident_date)
    add_conclusion(doc, failures)
    add_approvals(doc)
    timestamped_docx = ROOT / f"LMS_Test_Incident_Report_Yong_Di_Lun_NFR_Updated_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    saved_docx = None
    for candidate in [OUT_DOCX, FALLBACK_DOCX, timestamped_docx]:
        try:
            doc.save(candidate)
            saved_docx = candidate
            break
        except PermissionError:
            continue
    if saved_docx is None:
        raise PermissionError("Unable to save the Test Incident Report DOCX because all candidate files are locked.")
    build_markdown(data, specs)
    print(f"Created {saved_docx}")
    print(f"Created {OUT_MD}")
    print(f"Incidents: {len(failures)}")


if __name__ == "__main__":
    main()
