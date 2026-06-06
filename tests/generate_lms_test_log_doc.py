import json
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS_PATH = ROOT / "LMS_Test_Log_Results.json"
OUT_DOCX = ROOT / "LMS_Test_Log_Yong_Di_Lun_Executable.docx"
OUT_MD = ROOT / "LMS_Test_Log_Yong_Di_Lun_Executable.md"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, size=8, align=None, fill=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    for idx, line in enumerate(str(text).split("\n")):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Calibri"
        run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def fmt_date(iso_value):
    return datetime.fromisoformat(iso_value).strftime("%d/%m/%Y")


def fmt_time(iso_value):
    return datetime.fromisoformat(iso_value).strftime("%H%M")


def add_title(doc):
    title = doc.add_heading("Test Log", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_general_info(doc, data):
    table = doc.add_table(rows=13, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(2.7)

    header = table.cell(0, 0).merge(table.cell(0, 3))
    set_cell(header, "General Information", bold=True, fill="A6A6A6")

    set_cell(table.cell(1, 0), "Test Log Scope", fill="D9D9D9")
    set_cell(
        table.cell(1, 1).merge(table.cell(1, 3)),
        "This Test Log covered Library Management System features F001 to F0012 as described in Test Plan, LMS_TP_1.0.0.",
    )

    set_cell(table.cell(2, 0), "Test Log Description", fill="D9D9D9")
    set_cell(
        table.cell(2, 1).merge(table.cell(2, 3)),
        "The items tested were the LMS functional features. This test log records the execution result of test procedures in LMS_TPS_1.0.0.",
    )

    set_cell(table.cell(3, 0), "Version Author", fill="D9D9D9")
    set_cell(table.cell(3, 1), data["tester"])
    set_cell(table.cell(3, 2), "Contact Number", fill="D9D9D9")
    set_cell(table.cell(3, 3), "-")

    set_cell(table.cell(4, 0), "Revision Version", fill="D9D9D9")
    set_cell(table.cell(4, 1), "1.0")
    set_cell(table.cell(4, 2), "People Responsible", fill="D9D9D9")
    set_cell(table.cell(4, 3), "Tester 2")
    set_cell(table.cell(5, 0), "")
    set_cell(table.cell(5, 1), "")
    set_cell(table.cell(5, 2), "")
    set_cell(table.cell(5, 3), "Tester 3")
    set_cell(table.cell(6, 0), "")
    set_cell(table.cell(6, 1), "")
    set_cell(table.cell(6, 2), "")
    set_cell(table.cell(6, 3), data["tester"])

    header = table.cell(7, 0).merge(table.cell(7, 3))
    set_cell(header, "Activities Execution Information", bold=True, fill="A6A6A6")

    set_cell(table.cell(8, 0), "Execution Start Date", fill="D9D9D9")
    set_cell(table.cell(8, 1), fmt_date(data["start_time"]))
    set_cell(table.cell(8, 2), "End Date", fill="D9D9D9")
    set_cell(table.cell(8, 3), fmt_date(data["end_time"]))

    set_cell(table.cell(9, 0), "Execution Start Time", fill="D9D9D9")
    set_cell(table.cell(9, 1), fmt_time(data["start_time"]))
    set_cell(table.cell(9, 2), "End Time", fill="D9D9D9")
    set_cell(table.cell(9, 3), fmt_time(data["end_time"]))

    set_cell(table.cell(10, 0), "Tester Name", fill="D9D9D9")
    set_cell(table.cell(10, 1).merge(table.cell(10, 3)), data["tester"])

    set_cell(table.cell(11, 0), "Participant", fill="D9D9D9")
    set_cell(table.cell(11, 1).merge(table.cell(11, 3)), "-")

    set_cell(table.cell(12, 0), "Execution Summary", fill="D9D9D9")
    summary = data["summary"]
    set_cell(
        table.cell(12, 1).merge(table.cell(12, 3)),
        f"Total: {summary['total']}; Pass: {summary['pass']}; Fail: {summary['fail']}; Not Executed: {summary['not_executed']}",
    )
    doc.add_paragraph("")


def add_results(doc, rows):
    doc.add_paragraph("")
    table = doc.add_table(rows=2, cols=9)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    section = table.cell(0, 0).merge(table.cell(0, 8))
    set_cell(section, "Procedure Result", bold=True, fill="A6A6A6")

    headers = [
        "Requirement\nID",
        "Test Design\nID",
        "Test Case\nID",
        "Test Procedure\nID",
        "Type of\nTesting",
        "Tool",
        "Pass/Fail",
        "Test Incident Report ID/Test\nIncident ID",
        "Remark",
    ]
    widths = [0.85, 0.9, 0.75, 0.9, 1.05, 1.05, 0.75, 1.8, 2.8]
    for idx, header in enumerate(headers):
        set_cell(table.cell(1, idx), header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9D9D9")
        table.cell(1, idx).width = Inches(widths[idx])

    for result in rows:
        cells = table.add_row().cells
        values = [
            result["requirement_id"],
            result["test_design_id"],
            result["test_case_id"],
            result["test_procedure_id"],
            result["type_of_testing"],
            result["tool"],
            result["pass_fail"],
            result["incident_id"],
            result["remark"],
        ]
        for idx, value in enumerate(values):
            set_cell(cells[idx], value, size=7)
            cells[idx].width = Inches(widths[idx])
    doc.add_paragraph("")


def add_environment(doc, data):
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    section = table.cell(0, 0).merge(table.cell(0, 1))
    set_cell(section, "Environment Information", bold=True, fill="A6A6A6")
    set_cell(table.cell(1, 0), "Requested Test Environment", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9D9D9")
    set_cell(table.cell(1, 1), "Test Environment After Changes", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9D9D9")
    requested = (
        "Hardware: Local Windows workstation\n"
        "Software: Python 3.13.5, Selenium 4, Chrome headless, Python Requests/Postman-compatible HTTP automation\n"
        f"Application URL: {data['base_url']}\n"
        "Database: MySQL lms database\n"
        "Reference: LMS_TPS_1.0.0"
    )
    after = "No application logic was changed. Temporary test records were created and cleaned up by the test runner."
    set_cell(table.cell(2, 0), requested)
    set_cell(table.cell(2, 1), after)
    doc.add_paragraph("")


def add_anomalies(doc, rows, reporter):
    failures = [row for row in rows if row["pass_fail"] == "Fail"]
    not_executed = [row for row in rows if row["pass_fail"] == "Not Executed"]
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    section = table.cell(0, 0).merge(table.cell(0, 2))
    set_cell(section, "Anomalous Event", bold=True, fill="A6A6A6")
    headers = ["Unexpected event occurred", "Test Procedure ID", "Anomaly Reporter Name"]
    for idx, header in enumerate(headers):
        set_cell(table.cell(1, idx), header, bold=True, fill="D9D9D9")

    for row in failures + not_executed:
        cells = table.add_row().cells
        if row["pass_fail"] == "Not Executed":
            message = f"Not executed: {row['remark']}"
        else:
            message = row["remark"]
        set_cell(cells[0], message)
        set_cell(cells[1], row["test_procedure_id"])
        set_cell(cells[2], reporter)


def build_markdown(data):
    lines = [
        "# Test Log",
        "",
        f"Tester: {data['tester']}",
        f"Execution: {data['start_time']} to {data['end_time']}",
        f"Summary: Total {data['summary']['total']}, Pass {data['summary']['pass']}, Fail {data['summary']['fail']}, Not Executed {data['summary']['not_executed']}",
        "",
        "## Procedure Result",
        "",
        "| Requirement ID | Test Design ID | Test Case ID | Test Procedure ID | Type of Testing | Tool | Pass/Fail | Incident ID | Remark |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for row in data["results"]:
        lines.append(
            f"| {row['requirement_id']} | {row['test_design_id']} | {row['test_case_id']} | {row['test_procedure_id']} | "
            f"{row['type_of_testing']} | {row['tool']} | {row['pass_fail']} | {row['incident_id']} | {row['remark']} |"
        )
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main():
    data = json.loads(RESULTS_PATH.read_text(encoding="utf-8"), object_pairs_hook=OrderedDict)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(8)

    add_title(doc)
    add_general_info(doc, data)
    add_results(doc, data["results"])
    add_environment(doc, data)
    add_anomalies(doc, data["results"], data["tester"])
    doc.save(OUT_DOCX)
    build_markdown(data)
    print(f"Created {OUT_DOCX}")
    print(f"Created {OUT_MD}")


if __name__ == "__main__":
    main()
